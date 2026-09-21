from pathlib import Path

import pymupdf
import pytest
from docx import Document
from openpyxl import Workbook

from simulation_platform.sources import Packet, Source, read_packet, read_source


# Real bug found live, confirmed pre-existing in the original
# `simulation_platform` too (this file was already broken there the same
# way -- `test cases/` never actually lived at `.parents[1]` in either
# layout): the golden datasets live at the repo root
# (`Hackathon Specs/test cases/`), two levels up from this file
# (`<repo>/local-simulation-agent/tests/test_sources.py`), not one.
ROOT = Path(__file__).resolve().parents[2]


@pytest.mark.parametrize("name,count", [("tank", 12), ("iaq", 14), ("magnetic_circuit", 15), ("nacl_evaporation", 15)])
def test_every_document_in_real_packets_is_preserved(name, count):
    root = ROOT / "test cases" / (name + "_sysmlv2_full_dataset")
    packet = read_packet(root)
    assert len(packet.sources) == count
    assert all(source.path in packet.render() for source in packet.sources)
    assert sum(len(source.images) for source in packet.sources) > 0
    assert len(packet.render()) < 100000


def test_approved_and_superseded_values_survive_in_same_context():
    packet = read_packet(ROOT / "test cases" / "tank_sysmlv2_full_dataset")
    context = packet.render()
    assert "0.78" in context and "0.80" in context
    assert "CR-004 is now approved" in context
    assert "post-transfer" in context and "freeze" in context
    # Large reference trajectories cannot fill the generation context.
    reference = next(source for source in packet.sources if source.path.endswith(".csv"))
    assert reference.text not in context


def test_whole_input_schedule_is_available_and_reference_values_are_withheld():
    packet = Packet([Source("schedule.csv", "original rows", table=[{"x": str(i)} for i in range(101)]),
                     Source("expected.csv", "secret trajectory", table=[{"y": "999"}])])
    assert "original rows" not in packet.render()
    context = packet.render({"schedule.csv"})
    assert "original rows" in context
    assert "secret trajectory" not in context
    assert "999" not in context


def test_external_holdout_cannot_be_exposed_by_an_incorrect_model_classification():
    packet = Packet([Source("reference.csv", "time,y\n0,999\n", table=[{"time": "0", "y": "999"}])], {"reference.csv"})
    assert "999" not in packet.render({"reference.csv"})
    assert "Externally designated reference" in packet.render()


def test_context_limit_does_not_silently_drop_sources():
    with pytest.raises(ValueError, match="silently truncated"):
        Packet([Source("one.txt", "x" * 100)]).render(max_chars=50)


def test_scanned_pdf_is_sent_as_visual_evidence(tmp_path):
    source = pymupdf.open()
    page = source.new_page()
    page.insert_text((30, 30), "Flow = 0.006 m3/s")
    raster = page.get_pixmap().tobytes("png")
    scanned = pymupdf.open()
    scanned.new_page().insert_image(page.rect, stream=raster)
    path = tmp_path / "scan.pdf"
    scanned.save(path)
    scanned.close()
    source.close()
    result = read_source(path)
    assert len(result.images) == 1
    assert result.images[0].startswith("data:image/png;base64,")
    content = Packet([result]).content()
    assert any(block["type"] == "input_image" for block in content)


def test_excel_keeps_row_headers_units_and_formula_uncertainty(tmp_path):
    book = Workbook()
    sheet = book.active
    sheet.title = "Released"
    sheet.append(["Property", "Value", "Units"])
    sheet.append(["Area", 1.2, "m2"])
    sheet.append(["Volume", "=B2*2", "m3"])
    path = tmp_path / "register.xlsx"
    book.save(path)
    text = read_source(path).text
    assert "A1=Property | B1=Value | C1=Units" in text
    assert "A2=Area | B2=1.2 | C2=m2" in text
    assert "=B2*2" in text and "cached value=None" in text


def test_docx_retains_paragraph_table_order(tmp_path):
    doc = Document()
    doc.add_paragraph("The following row is superseded:")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Old limit", "0.78"
    doc.add_paragraph("Approved replacement is 0.80 m.")
    path = tmp_path / "review.docx"
    doc.save(path)
    text = read_source(path).text
    assert text.index("superseded") < text.index("0.78") < text.index("0.80")


def test_unsupported_or_corrupt_inputs_stop_with_exact_file(tmp_path):
    (tmp_path / "unreadable.xyz").write_bytes(b"unsupported")
    with pytest.raises(ValueError, match="unreadable.xyz"):
        read_packet(tmp_path)
