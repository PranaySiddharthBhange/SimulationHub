"""DOCX parser. Mirrors `Document Agent.md`, Section 9 (DOCX row)."""

from __future__ import annotations

from pathlib import Path

import docx

from engineering_agent.schemas import Observation, ObservationType, SourceLocation

from .base import BaseParser, ParserError


class DocxParser(BaseParser):
    parser_name = "docx_parser"
    extensions = (".docx",)

    def parse(self, document, file_path: Path) -> list[Observation]:
        try:
            doc = docx.Document(str(file_path))
        except Exception as exc:
            raise ParserError(f"Unable to open DOCX: {exc}") from exc

        observations: list[Observation] = []
        current_section = "(document start)"

        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                continue
            if para.style is not None and para.style.name.startswith("Heading"):
                current_section = text
            observations.append(
                Observation(
                    observation_id=self.next_observation_id(document),
                    document_id=document.document_id,
                    type=ObservationType.TEXT,
                    content=text,
                    location=SourceLocation(section=current_section),
                    parser=self.parser_name,
                    parser_version=self.parser_version,
                )
            )

        for table_index, table in enumerate(doc.tables, start=1):
            rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
            if rows:
                observations.append(
                    Observation(
                        observation_id=self.next_observation_id(document),
                        document_id=document.document_id,
                        type=ObservationType.TABLE,
                        content="\n".join(" | ".join(row) for row in rows),
                        location=SourceLocation(section=f"Table {table_index}"),
                        parser=self.parser_name,
                        parser_version=self.parser_version,
                    )
                )

        if not observations:
            raise ParserError("DOCX produced no text or tables.")
        return observations
